from docx import Document
import re

file_list = [
    [2013, "1.0", "//dkg-dc-01/Daten/Daten/Arbeitsverzeichnisse/OL/1 Dokumente zu laufenden LL/HCC/1. Förderung (2009-2013) Version 1/Version 1/LL Versionen/S3LLHCC_Kurzversion_20130226.docx"],
    [2021, "2.0", "//dkg-dc-01/Daten/Daten/Arbeitsverzeichnisse/OL/1 Dokumente zu laufenden LL/HCC/2. Förderung (2018-2021) Version 2 und 3/Version 2 (ab 2018)/Leitlinienversionen/Kurzversion/LL_HCC&biliäre_Karzinome_Kurzversion 2.0_.docx"],
    [2022, "3.1", "//dkg-dc-01/Daten/Daten/Arbeitsverzeichnisse/OL/1 Dokumente zu laufenden LL/HCC/2. Förderung (2018-2021) Version 2 und 3/Version 3/Leitlinienversionen/Kurzversion/LL_Hepatozelluläres Karzinom und biliäre Karzinome_Kurzversion_3.1.docx"],
    [2023, "4.0", "//dkg-dc-01/Daten/Daten/Arbeitsverzeichnisse/OL/1 Dokumente zu laufenden LL/HCC/3. Förderung (2022-20  ) Version 4, 5/Leitlinienversionen/Version 4/Leitlinienversionen/Kurzversion/LL_Hepatozelluläres Karzinom und biliäre Karzinome_Kurzversion_4.0_24.08.docx"],
]

out_csv = "Leitlinie|Version|Jahr|Fullname|Kapitel 1|Kapitel 2|Kapitel 3|Thema|Nummer|Text|Typ|LoE|GoR\n"


def split_into_words(text):
    # Define your stop symbols
    stop_symbols = ['.', ',', ';', ':', '!', '?', '"', "'", '-', '(', ')', '[', ']', '{', '}', '/', '\\', '|', '_', '+',
                    '=', '*', '&', '^', '%', '$', '#', '@', '~', '`']

    # Split the text into words
    words = re.split(r'\s+', text)

    # Filter out stop symbols
    words = [word for word in words if word.strip() and word not in stop_symbols]

    return words

def detect_merged_cells(row):
    merged_cells = set()
    for cell in row.cells:
        if cell._element.grid_span > 1 or cell._element.vMerge:
            merged_cells.add((cell._element.getparent().index(cell._element), cell._element.grid_span))
    return merged_cells

for file in file_list:
    doc = Document(file[2])
    for table in doc.tables:
        first_cell = table.cell(0,0)
        if first_cell.text != "Nr.":
            continue

        for row_idx in range(1, len(table.rows)):
            row = table.rows[row_idx]
            reco_number = row.cells[0].text.replace(".", "_").strip("_")
            reco_text = row.cells[1].text.replace("\r\n", " ").replace("\r", " ").replace("\n", " ").strip()
            column_count = len(row.cells)

            if (row.cells[2].text == row.cells[3].text and row.cells[2].text == row.cells[4].text) or len([x for x in row.cells if len(x.text) == 0]) > 1:
                reco_type = "Konsensbasiert"
                gor_list = []
                text_arr = split_into_words(reco_text)
                if any(x for x in text_arr if x == "können" or x == "kann"):
                    gor_list.append("0")
                if any(x for x in text_arr if x == "sollte" or x == "sollten"):
                    gor_list.append("B")
                if any(x for x in text_arr if x == "soll" or x == "sollen"):
                    gor_list.append("A")
                reco_gor = ",".join(gor_list)
                reco_loe = "N/A"
            else:
                reco_loe = row.cells[3].text
                if row.cells[2].text == "ST":
                    reco_type = "Evidenzbasiertes Statement"
                    reco_gor = "N/A"
                else:
                    reco_type = "Evidenzbasierte Empfehlung"
                    reco_gor = row.cells[2].text.replace("\r\n", " ").replace("\r", " ").replace("\n", " ").strip()
            pass
            new_line = "%s|%s|%s|%s|%s|%s|%s|%s|%s|%s|%s|%s|%s" % ("HCC", file[1], file[0], "%s V%s (%s)" % ("HCC", file[1], file[0]), "", "", "", "", reco_number, reco_text, reco_type, reco_loe, reco_gor)
            out_csv += new_line.replace("\r\n", "").replace("\r", "").replace("\n", "").strip()+"\n"
        pass

with open("HCC_recos.csv", "w", encoding="utf-8") as f:
    f.write(out_csv)